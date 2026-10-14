"""
NSE All Stock List Fetcher using NseIndiaApi
Fetches all NSE equity symbols from the bhavcopy and saves them to a file.
Handles NSE's session requirements automatically
"""

import os
from datetime import datetime, timedelta
from pathlib import Path

from nse import NSE
import pandas as pd


def get_all_equity_symbols(nse: NSE) -> list[str]:
    """Fetch all equity symbols from NSE bhavcopy."""
    for days_back in [0, 1, 2, 3, 4, 5]:
        try:
            target_date = datetime.now() - timedelta(days=days_back)
            path = nse.equityBhavcopy(target_date)
            df = pd.read_csv(path)
            # Filter for equity series (EQ) and stock type (STK)
            eq = df[(df['SctySrs'] == 'EQ') & (df['FinInstrmTp'] == 'STK')]
            symbols = eq['TckrSymb'].tolist()
            if symbols:
                print(f"Loaded {len(symbols)} equity symbols from bhavcopy ({target_date})")
                # Clean up old bhavcopy file
                if os.path.exists(path):
                    os.remove(path)
                    print(f"Deleted old bhavcopy: {path}")
                return symbols
        except Exception as e:
            print(f"Bhavcopy not available for {target_date}: {e}")
            continue
    print("Warning: Could not fetch bhavcopy")
    return []


def save_stocks_to_file(symbols: list[str], filename: str = "nse_stocks.txt") -> None:
    """Save the list of stock symbols to a file."""
    with open(filename, "w") as f:
        for symbol in symbols:
            f.write(f"{symbol}\n")
    print(f"Saved {len(symbols)} symbols to {filename}")


def main():
    print("\nNSE All Stock List Fetcher")
    with NSE(download_folder=Path.cwd(), timeout=20) as nse:
        stocks = get_all_equity_symbols(nse)
        if stocks:
            save_stocks_to_file(stocks)
        else:
            print("No symbols fetched.")


if __name__ == "__main__":
    main()


import re
import openpyxl

# Load nse_stocks.txt
with open('nse_stocks.txt', 'r') as f:
    nse_stocks = [line.strip() for line in f if line.strip()]

# Load stocks from Stcapmark.xlsx
wb = openpyxl.load_workbook('Stcapmark.xlsx', read_only=True)
ws = wb.active
stcap_stocks = []
for row in ws.iter_rows(min_row=2, max_col=1):
    val = row[0].value
    if val:
        stcap_stocks.append(val)
wb.close()

# Load stocks from apinse.py (fallback list)
with open('apinse.py', 'r') as f:
    content = f.read()
match = re.search(r'stocks \+= \[(.*?)\]', content, re.DOTALL)
apinse_stocks = re.findall(r'"([^"]+)"', match.group(1)) if match else []

nse_set = set(nse_stocks)
stcap_set = set(stcap_stocks)
apinse_set = set(apinse_stocks)

# Find stocks in nse_stocks.txt NOT in Stcapmark.xlsx
missing_from_stcap = nse_set - stcap_set
print(f"NSE stocks: {len(nse_stocks)}")
print(f"Stcapmark.xlsx stocks: {len(stcap_stocks)}")
print(f"apinse.py stocks (fallback): {len(apinse_stocks)}")
print(f"Missing from Stcapmark.xlsx: {len(missing_from_stcap)}")
print(f"Missing from apinse.py: {len(nse_set - apinse_set)}")

# Write missing stocks to files for review
with open('missing_from_stcap.txt', 'w') as f:
    for s in sorted(missing_from_stcap):
        f.write(s + chr(10))

with open('missing_from_apeinse.txt', 'w') as f:
    for s in sorted(nse_set - apinse_set):
        f.write(s + chr(10))


import re
import openpyxl
from pathlib import Path

# Load nse_stocks.txt
with open('nse_stocks.txt', 'r') as f:
    nse_stocks = [line.strip() for line in f if line.strip()]
nse_set = set(nse_stocks)

# --- 1. Add missing stocks to Stcapmark.xlsx ---
wb = openpyxl.load_workbook('Stcapmark.xlsx')
ws = wb.active

# Collect existing stocks from column A
existing_stocks = set()
for row in range(2, ws.max_row + 1):
    val = ws.cell(row=row, column=1).value
    if val:
        existing_stocks.add(val)

missing_from_stcap = nse_set - existing_stocks
print(f"Missing from Stcapmark.xlsx: {sorted(missing_from_stcap)}")

for stock in sorted(missing_from_stcap):
    ws.cell(row=ws.max_row + 1, column=1, value=stock)
    print(f"  Added '{stock}' to Stcapmark.xlsx")

wb.save('Stcapmark.xlsx')
wb.close()
print(f"Stcapmark.xlsx saved. Total rows: {ws.max_row}")

# --- 2. Add missing stocks to apinse.py ---
with open('apinse.py', 'r') as f:
    content = f.read()

match = re.search(r'stocks \+= \[(.*?)\]', content, re.DOTALL)
if not match:
    print("Could not find stocks list in apinse.py")
    exit(1)

existing_apeinse = set(re.findall(r'"([^"]+)"', match.group(1)))
missing_from_apeinse = nse_set - existing_apeinse
print(f"Missing from apinse.py: {len(missing_from_apeinse)}")

if missing_from_apeinse:
    # Find the closing bracket position of the stocks list
    close_bracket_pos = content.index(']', match.start(0))

    # Get the content just before the closing bracket (strip whitespace)
    before_bracket = content[:close_bracket_pos].rstrip()

    # Ensure the last existing stock has a trailing comma
    if not before_bracket.endswith(','):
        before_bracket += ','

    # Build new entries: comma-separated, last one has NO trailing comma
    new_entries = '\n    ' + ',\n    '.join(f'"{s}"' for s in sorted(missing_from_apeinse)) + '\n'

    # Rebuild content: existing content + comma + new entries + closing bracket
    new_content = before_bracket + new_entries + content[close_bracket_pos:]

    with open('apinse.py', 'w') as f:
        f.write(new_content)

    print(f"apinse.py updated. Added {len(missing_from_apeinse)} stocks.")
else:
    print("apinse.py already up to date.")
import re
import openpyxl

# Verify Stcapmark.xlsx
wb = openpyxl.load_workbook('Stcapmark.xlsx', read_only=True)
ws = wb.active
stcap_stocks = []
for row in ws.iter_rows(min_row=2, max_col=1):
    val = row[0].value
    if val:
        stcap_stocks.append(val)
wb.close()

print(f"Stcapmark.xlsx total stocks: {len(stcap_stocks)}")
print(f"RELIANCE in Stcapmark.xlsx: {'RELIANCE' in stcap_stocks}")

# Verify apinse.py
with open('apinse.py', 'r') as f:
    content = f.read()
match = re.search(r'stocks \+= \[(.*?)\]', content, re.DOTALL)
sym_list = re.findall(r'"([^"]+)"', match.group(1))
print(f"apinse.py total stocks: {len(sym_list)}")
print(f"20MICRONS in apinse.py: {'20MICRONS' in sym_list}")
print(f"RELIANCE in apinse.py: {'RELIANCE' in sym_list}")
print(f"Last 3 stocks in apinse.py: {sym_list[-3:]}")

# Cross-check
with open('nse_stocks.txt', 'r') as f:
    nse_stocks = [line.strip() for line in f if line.strip()]

nse_set = set(nse_stocks)
stcap_set = set(stcap_stocks)
sym_set = set(sym_list)

still_missing_stcap = nse_set - stcap_set
still_missing_apeinse = nse_set - sym_set

print(f"\nStill missing from Stcapmark.xlsx: {len(still_missing_stcap)}")
print(f"Still missing from apinse.py: {len(still_missing_apeinse)}")
